import streamlit as st
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import io
import json
from google import genai
from google.genai import types

# إعدادات واجهة البرنامج
st.set_page_config(page_title="نظام تفريغ لوحات السيارات", layout="centered")

# تطبيق اتجاه اليمين لليسار في واجهة Streamlit
st.markdown("""
<style>
    body { direction: rtl; text-align: right; }
    .stApp { direction: rtl; text-align: right; }
    .stMarkdown, .stText, p, h1, h2, h3 { text-align: right; }
</style>
""", unsafe_allow_html=True)

st.title("🚗 نظام تفريغ وسحب بيانات لوحات السيارات")
st.write("ارفع التسجيلات الصوتية لتفريغها مباشرة وفق الضوابط المحددة وتصديرها كملف Excel.")

# مفتاح API
api_key = st.text_input("أدخل مفتاح Gemini API Key الخاص بك:", type="password")

# التعليمات والضوابط الصارمة
SYSTEM_INSTRUCTIONS = """
أنت نظام خبير في تفريغ لوحات السيارات والمواقع من المقاطع الصوتية بدقة 100%.
يجب عليك الالتزام بالضوابط والقواعد التالية بدقة متناهية ودون كتابة أي ملاحظات:

1. هيكل وجدول البيانات:
- العمود A (حروف وأرقام السيارة):
  * دمج حروف اللوحة معاً وإلغاء المسافات بينها.
  * تبديل أي ألف مهموزة (أ / إ / آ) إلى ألف قائمة/عادية (ا).
  * توحيد حرف الهاء بالشكل العادي (ه) دائماً بدلاً من (هـ).
  * إلحاق أرقام اللوحة مباشرة بعد الحروف المدمجة (بدون مسافات).
  * إلغاء المسافات بين الحروف في أول كولوم والمسافات بين الأرقام كذلك.
  * الحروف التالية مستحيل أن تكون موجودة في اللوحات (العمود A): (ج، خ، غ، ف، ث، ض، ة، ت، ش، ز، ذ). إذا سمعت أي منها، قم بتصحيحها بناءً على النطق الصوتي الصحيح المشابه للوحات النظامية (مثل ف->ق، غ->ع، ض->ص، ش->س، ت/ذ->د، ز->ر).

- العمود B (رقم الموقع):
  * كتابة رقم الموقع مرة واحدة فقط في أول سطر عند بداية كل موقع جديد.
  * ترك الخلية فارغة تماماً لباقي سيارات نفس الموقع إلى أن يُذكر موقع جديد.

- العمود C (التصنيف والملاحظات):
  * نقل ⬅️ يُكتب: ن
  * تاكسي ⬅️ يُكتب: ت
  * سجل عليها حرف الباء ⬅️ يُكتب: ب
  * سجل عليها حرف الميم ⬅️️ يُكتب: م
  * سجل عليها حرف الفاء ⬅️ يُكتب: ف
  * سجل عليها حرف الراء ⬅️ يُكتب: ر
  * مربع ⬅️ يُكتب: مربع
  * شقق ⬅️ يُكتب: شقق
  * الحالات المركبة (الدمج):
    نقل + حرف الباء ⬅️ يُكتب: ن ب
    نقل + حرف الميم ⬅️ يُكتب: ن م
    نقل + حرف الفاء ⬅️ يُكتب: ن ف
    تطبيق نفس نمط الدمج على أي ملاحظات متكررة أو مركبة بنفس الطريقة.
  * الخانات الفارغة: تكون فارغة كلياً وبشكل تام (خالية من أي مسافات أو رموز أو مساحات بيضاء).

2. الترتيب التسلسلي:
- مراعاة الترتيب التسلسلي الصحيح للمواقع واللوحات بدقة متناهية.

يجب إرجاع النتيجة بصيغة JSON فقط كقائمة من الكائنات (Array of objects)، حيث يحتوي كل كائن على:
- "plate": قيمة العمود A
- "site": قيمة العمود B (فارغة "" إذا لم تكن بداية موقع جديد)
- "note": قيمة العمود C (فارغة "" إذا لم توجد ملاحظة)
"""

uploaded_files = st.file_uploader(
    "اختر التسجيلات الصوتية (يمكنك اختيار عدة تسجيلات معاً):", 
    type=["m4a", "mp3", "wav", "aac", "ogg"], 
    accept_multiple_files=True
)

def create_excel_file(records):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "بيانات السيارات"
    
    # اتجاه ورقة العمل من اليمين إلى اليسار
    ws.views.sheetView[0].rightToLeft = True

    # العناوين
    headers = ["حروف وأرقام السيارة", "رقم الموقع", "التصنيف والملاحظات"]
    ws.append(headers)

    # التنسيقات
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    data_font = Font(name="Calibri", size=11)
    align_center = Alignment(horizontal="center", vertical="center")

    for col_idx, cell in enumerate(ws[1], 1):
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    # تعبئة البيانات
    for row_idx, item in enumerate(records, start=2):
        plate = str(item.get("plate", "")).strip()
        site = str(item.get("site", "")).strip() or None
        note = str(item.get("note", "")).strip() or None
        
        ws.append([plate, site, note])
        for col_idx in range(1, 4):
            c = ws.cell(row=row_idx, column=col_idx)
            c.font = data_font
            c.alignment = align_center
            c.border = thin_border

    # ضبط عرض الأعمدة تلقائياً
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 6, 18)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

if st.button("بدء التفريغ واستخراج الإكسيل", type="primary"):
    if not api_key:
        st.error("يرجى إدخال مفتاح API أولاً.")
    elif not uploaded_files:
        st.warning("يرجى رفع ملف صوتي واحد على الأقل.")
    else:
        try:
            client = genai.Client(api_key=api_key)
            all_records = []
            
            with st.spinner("جاري معالجة التسجيلات الصوتية عبر جيميناي بدقة..."):
                for uploaded_file in uploaded_files:
                    st.write(f"معالجة الملف: `{uploaded_file.name}`...")
                    
                    audio_bytes = uploaded_file.read()
                    mime_type = uploaded_file.type or "audio/m4a"

                    response = client.models.generate_content(
                        model='gemini-2.5-flash',
                        contents=[
                            types.Part.from_bytes(
                                data=audio_bytes,
                                mime_type=mime_type,
                            ),
                            "قم بتفريغ هذا التسجيل الصوتي بدقة تامة طبقاً لتعليمات النظام بدون أي تعديل أو ملاحظات إضافية."
                        ],
                        config=types.GenerateContentConfig(
                            system_instruction=SYSTEM_INSTRUCTIONS,
                            response_mime_type="application/json"
                        )
                    )
                    
                    data = json.loads(response.text)
                    if isinstance(data, list):
                        all_records.extend(data)

            if all_records:
                st.success(f"تم تفريغ البيانات بنجاح! إجمالي السجلات: {len(all_records)}")
                
                # عرض معاينة
                df_preview = pd.DataFrame(all_records)
                df_preview.columns = ["حروف وأرقام السيارة", "رقم الموقع", "التصنيف والملاحظات"]
                st.dataframe(df_preview, use_container_width=True)

                # إنشاء وتحميل الملف
                excel_data = create_excel_file(all_records)
                st.download_button(
                    label="📥 تحميل ملف الإكسيل الجاهز (.xlsx)",
                    data=excel_data,
                    file_name="تفريغ_بيانات_السيارات.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
            else:
                st.warning("لم يتم استخراج أي سجلات من الملفات.")

        except Exception as e:
            st.error(f"حدث خطأ أثناء المعالجة: {str(e)}")
